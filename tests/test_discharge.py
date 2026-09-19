"""Acute discharge sitrep ingest, on SYNTHETIC fixtures.

tests/fixtures/discharge_*.xlsx come from make_external_fixtures.py: invented trusts and
counts in the two Table 2 shapes found in the real files (the first 2022 publications,
four columns per day with discharges split at 17:00 and helper rows of item codes; and
everything since the 2023 re-issues, three columns per day under an ICB section).
"""
from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from nhs_ae.ingest import discharge, external

FIX = Path(__file__).parent / "fixtures"
SPLIT = FIX / "discharge_2022_split.xlsx"
SINGLE = FIX / "discharge_2024_single.xlsx"
UP = "https://www.england.nhs.uk/statistics/wp-content/uploads/sites/2/"


# --------------------------------------------------------------------------------------
# Links and specifications
# --------------------------------------------------------------------------------------
@pytest.mark.parametrize("href, text, period", [
    ("2023/12/Daily-discharge-sitrep-monthly-data-webfile-April2022-revised.xlsx",
     "Daily-discharge-sitrep-monthly-data-webfile-April2022 revised", date(2022, 4, 1)),
    ("2022/07/Daily-discharge-sitrep-monthly-data-webfile-June2022.xlsx", "", date(2022, 6, 1)),
    ("2026/05/Daily-discharge-sitrep-monthly-data-webfile-01-April2026-v1.xlsx", "",
     date(2026, 4, 1)),
    ("2026/04/Daily-discharge-sitrep-monthly-data-webfile-January-2026.xlsx", "",
     date(2026, 1, 1)),
    ("2024/10/Daily-discharge-sitrep-monthly-data-webfile-September20241-2.xlsx", "",
     date(2024, 9, 1)),
    ("2025/10/Daily-discharge-sitrep-monthly-data-webfile-September2024v2.xlsx", "",
     date(2024, 9, 1)),
])
def test_parse_link(href, text, period):
    link = discharge.parse_link(UP + href, text)
    assert link is not None and link.period == period


@pytest.mark.parametrize("href", [
    "2026/09/Daily-discharge-sitrep-monthly-data-webfile-05-CSV-aug26.csv",
    "2024/11/Daily-discharge-sitrep-monthly-data-webfile-October2024-csv.xlsx",
    "2026/09/Daily-discharge-sitrep-timeseries-data-webfile-April2021-August2026.xlsx",
    "2023/11/Community-discharge-sitrep-monthly-data-webfile-October2023.xlsx",
    "2023/11/Discharge-Ready-Date-monthly-data-webfile-September2023.xlsx",
    "2021/06/Daily-discharge-sitrep-monthly-data-webfile-May2021.xlsx",  # before April 2022
])
def test_parse_link_rejects_other_files(href):
    assert discharge.parse_link(UP + href, "") is None


@pytest.mark.parametrize("period, collection, changed", [
    (date(2022, 4, 1), "covid_eprr_discharge_sitrep", False),
    (date(2022, 5, 1), "covid_eprr_discharge_sitrep", True),
    (date(2022, 6, 1), "adsr_spec_2022", False),
    (date(2024, 4, 1), "adsr_spec_2022", False),
    (date(2024, 5, 1), "adsr_spec_2022", True),
    (date(2024, 6, 1), "adsr_spec_2024", False),
    (date(2026, 8, 1), "adsr_spec_2024", False),
])
def test_collection_for(period, collection, changed):
    assert discharge.collection_for(period) == (collection, changed)


@pytest.mark.parametrize("header, measure", [
    ("Number of patients who no longer meet the criteria to reside", "nctr"),
    ("Number of patients remaining in hospital who no longer meet the criteria to reside",
     "remaining"),
    ("Number of patients discharged", "discharged"),
    ("Number of patients discharged by 17:00", "discharged_by_17"),
    ("Number of patients discharged between 17:01 and 23:59", "discharged_after_17"),
    ("DIS002_TOTAL", None), (None, None),
])
def test_classify_measure(header, measure):
    assert discharge.classify_measure(header) == measure


# --------------------------------------------------------------------------------------
# Workbook layouts
# --------------------------------------------------------------------------------------
def test_read_2022_split_layout():
    f = discharge.read_discharge(SPLIT)
    assert (f.period_start, f.period_end) == (date(2022, 7, 1), date(2022, 7, 31))
    assert f.published == date(2022, 8, 11) and f.revised is None  # "Revised: -"
    assert f.layout == "t2_4col|split_discharges|3tables"
    rows = f.rows.set_index("org_code")
    assert list(rows.index) == ["ENG", "ZZA", "ZZB"]  # header row and notes are not trusts
    assert rows.loc["ENG", "org_name"] == "ENGLAND (All Acute Trusts)"
    zza = rows.loc["ZZA"]
    assert zza.nctr_avg == pytest.approx(100.0)
    assert zza.discharged_avg == pytest.approx((30 + 30 + 30 + 40) / 4)  # by 17:00 + after
    assert zza.remaining_avg == pytest.approx(zza.nctr_avg - zza.discharged_avg)
    assert zza.region == "EAST OF ENGLAND"
    zzb = rows.loc["ZZB"]
    assert zzb.days_reported == 3 and zzb.days_in_period == 4  # '-' on one day
    assert zzb.nctr_avg == pytest.approx(60.0)


def test_read_2024_single_layout_skips_icb_rows():
    f = discharge.read_discharge(SINGLE.read_bytes())
    assert f.published == date(2024, 9, 12) and f.revised is None
    assert f.layout == "t2_3col|single_discharges|7tables"
    rows = f.rows.set_index("org_code")
    assert "ZZ9" not in rows.index  # the ICB row sits above the trust header
    assert rows.loc["ENG", "org_name"] == "ENGLAND (Type 1 Trusts)"
    # England is published, not derived: trusts + 5 a day (the hospice rule), ZZB's
    # missing day counted as zero -> (155 + 115 + 165 + 165) / 4
    assert rows.loc["ENG", "nctr_avg"] == pytest.approx(150.0)
    assert rows.loc["ZZA", "discharged_avg"] == pytest.approx(32.5)
    assert rows.loc["ZZB", "days_reported"] == 3


def test_missing_table_2_is_an_error(tmp_path):
    wb = load_workbook(SINGLE)
    del wb["Table 2"]
    p = tmp_path / "x.xlsx"
    wb.save(p)
    with pytest.raises(ValueError, match="Table 2"):
        discharge.read_discharge(p)


def test_read_dates_from_cover_sheet():
    assert discharge.read_dates(SPLIT.read_bytes()) == (date(2022, 8, 11), None)


def test_workbook_without_cover_sheet(tmp_path):
    """October 2024 was published with no cover sheet: no dates, period from Table 2."""
    wb = load_workbook(SINGLE)
    del wb["Cover Sheet"]
    p = tmp_path / "nocover.xlsx"
    wb.save(p)
    f = discharge.read_discharge(p)
    assert f.published is None and f.revised is None
    assert f.period_start == date(2024, 8, 1)
    assert discharge.read_dates(p.read_bytes()) == (None, None)


def test_zero_on_the_england_row_is_a_missing_day(tmp_path):
    """England shows 0 on 19 June 2022 (a lost day) where trusts show '-'."""
    wb = load_workbook(SINGLE)
    ws = wb["Table 2"]
    for row in ws.iter_rows():
        if row[3].value == "ENGLAND (Type 1 Trusts)":
            for c in row[7:10]:  # the second day's three columns
                c.value = 0
    p = tmp_path / "zero.xlsx"
    wb.save(p)
    eng = discharge.read_discharge(p).rows.set_index("org_code").loc["ENG"]
    assert eng.days_reported == 3
    assert eng.nctr_avg == pytest.approx((155 + 165 + 165) / 3)


def test_first_publication_falls_back_to_version_dates():
    assert external.first_publication([None, date(2024, 11, 14)], ["2025-01-21", None]) == \
        date(2024, 11, 14)
    assert external.first_publication([None], ["2025-01-21", "2025-02-14"]) == date(2025, 1, 21)
    assert external.first_publication([], []) is None


# --------------------------------------------------------------------------------------
# Build: versions and as-of
# --------------------------------------------------------------------------------------
def _variant(tmp_path: Path, name: str, revised, bump: float = 0.0) -> Path:
    """The 2024 fixture with a Revised date and every ZZA NCtR value shifted by ``bump``."""
    wb = load_workbook(SINGLE)
    wb["Cover Sheet"]["C8"] = revised
    ws = wb["Table 2"]
    for row in ws.iter_rows():
        if row[2].value == "ZZA":
            for c in row[4::3]:
                if isinstance(c.value, (int, float)):
                    c.value += bump
    p = tmp_path / "data" / "external" / "discharge" / "raw" / name
    p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(p)
    return p


def test_build_versions_and_as_of(tmp_path):
    orig = _variant(tmp_path, "orig.xlsx", None)
    rev = _variant(tmp_path, "rev.xlsx", date(2025, 4, 10), bump=8)
    manifest = tmp_path / "data" / "external" / "discharge" / "raw" / "manifest.jsonl"
    for p, pub, why in [(orig, "2024-09-12", "title_block_published"),
                        (rev, "2025-04-10", "title_block_revised")]:
        external.append_jsonl({"url": UP + p.name, "filename": p.name, "period": "2024-08-01",
                               "stored": True, "sha256": p.name,
                               "path": str(p.relative_to(tmp_path)), "published": pub,
                               "published_source": why}, manifest)
    df, failures = discharge.build(manifest, tmp_path, out_path=None)
    assert failures.empty
    assert (df.period == pd.Timestamp("2024-08-01")).all()
    assert (df.collection == "adsr_spec_2024").all() and not df.spec_change_in_month.any()
    assert (df.first_published == pd.Timestamp("2024-09-12")).all()
    zza = df[df.org_code == "ZZA"].set_index("version")
    assert zza.loc[2, "nctr_avg"] - zza.loc[1, "nctr_avg"] == pytest.approx(8)
    assert zza.loc[2, "is_latest"] and not zza.loc[1, "is_latest"]

    def nctr(when):
        cur = discharge.as_of(df, when)
        return cur.loc[cur.org_code == "ZZA", "nctr_avg"].tolist()

    assert nctr("2024-09-11") == []
    assert nctr("2024-10-01") == [zza.loc[1, "nctr_avg"]]
    assert nctr("2025-04-10") == [zza.loc[2, "nctr_avg"]]

    vt = discharge.versions(df).set_index("version")
    assert vt.loc[1, "lag_days"] == 12  # 12 September after 31 August
    assert vt.loc[2, "orgs_changed"] == 1
    assert vt.loc[2, "england_nctr_avg_change"] == 0  # England row untouched by the bump
    assert np.isnan(df.loc[df.org_code == "ZZB", "nctr_avg"]).sum() == 0


def test_build_uses_the_files_period_over_the_link(tmp_path, caplog):
    p = _variant(tmp_path, "mislabelled.xlsx", None)
    manifest = tmp_path / "m.jsonl"
    external.append_jsonl({"url": UP + p.name, "filename": p.name, "period": "2024-07-01",
                           "stored": True, "sha256": "s", "path": str(p.relative_to(tmp_path)),
                           "published": "2024-09-12"}, manifest)
    df, _ = discharge.build(manifest, tmp_path, out_path=None)
    assert (df.period == pd.Timestamp("2024-08-01")).all()
    assert "file says 2024-08-01" in caplog.text
