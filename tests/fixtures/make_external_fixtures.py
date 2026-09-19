"""Write the SYNTHETIC KH03 and discharge workbooks used by test_kh03.py / test_discharge.py.

Every number, organisation code and name here is invented. The files copy only the
*shape* of the real layouts (title blocks, header rows, merged-group rows, helper rows,
sheet order), which is what the readers depend on. Real KH03 files before 2013-14 Q4 are
.xls; the reader goes through pandas either way, so the fixtures are .xlsx (no xlwt).

    python tests/fixtures/make_external_fixtures.py
"""

from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).parent
SECTORS = ["Total ", "General & Acute", "Learning Disabilities", "Maternity", "Mental Illness"]


def _sheet(wb: Workbook, title: str, rows: list[list], first: bool = False):
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    for r in rows:
        ws.append(r)
    return ws


def _kh03_block(lead: list, name: str, avail: list, occ: list) -> list:
    """lead cells + Available(5) + gap + Occupied(5) + gap + % Occupied(5)."""
    pct = [("-" if a == 0 else round(o / a, 6)) for a, o in zip(avail, occ)]
    return lead + [name] + avail + [None] + occ + [None] + pct


def kh03_status_line() -> Path:
    """2010-11 to 2012-13: title in column A, one 'Status:' line, SHA Code, PCT rows."""
    wb = Workbook()
    # (lead cells, name, available by sector, occupied by sector)
    trust_rows = [
        (["2010-11", "June", "Q90", "ZZ1"], "Synthetic Teaching PCT",
         [12, 0, 12, 0, 0], [9.5, 0, 9.5, 0, 0]),
        (["2010-11", "June", "Q90", "ZZA"], "Synthetic Hospitals NHS Trust",
         [400, 360, 0, 40, 0], [340, 320, 0, 20, 0]),
        (["2010-11", "June", "Q91", "ZZB"], "Another Synthetic NHS Foundation Trust",
         [600, 500, 0, 50, 50], [520, 450, 0, 30, 40]),
    ]
    tot_a = [sum(r[2][i] for r in trust_rows) for i in range(5)]
    tot_o = [sum(r[3][i] for r in trust_rows) for i in range(5)]
    header = ["Year", "Period", "SHA Code", "Org Code", "Org Name"] + SECTORS + [None] + \
        SECTORS + [None] + SECTORS
    rows = [
        ["Title:", ("Average daily number of available and occupied beds open overnight by "
                    "sector, Quarter 1 2010-11 (SYNTHETIC)")],
        ["Source:", "Department of Health form KH03"],
        ["Status:", "Published 18 November 2010 and revised 24 May 2012"],
        [],
        [None] * 5 + ["Available"] + [None] * 5 + ["Occupied"] + [None] * 5 + ["% Occupied"],
        header,
        [],
        _kh03_block(["2010-11", "June", None, None], "England", tot_a, tot_o),
        [],
    ] + [_kh03_block(*r) for r in trust_rows] + [[], ["Footnote: synthetic data."]]
    _sheet(wb, "NHS Trust by Sector", rows, first=True)
    _sheet(wb, "SHA by Sector", [["Title:", "synthetic"]])
    _sheet(wb, "Occupied by Specialty", [["Title:", "synthetic"]])
    _sheet(wb, "Data quality", [
        ["Data quality statement"], [],
        [None, "The following organisations did not supply data for this quarter:"],
        [None, "Synthetic Mental Health NHS Trust"]])
    out = HERE / "kh03_status_line.xlsx"
    wb.save(out)
    return out


def _kh03_modern(path: Path, *, published, revised, parent: str, month_col: str, ld: str,
                 period_text: str, year: str, month: str, dq_first: bool,
                 dq_rows: list[list], trusts: list) -> Path:
    wb = Workbook()
    sectors = [s.replace("Learning Disabilities", ld) for s in SECTORS]
    header = [None, "Year", month_col, parent, "Org Code", "Org Name"] + sectors + [None] + \
        sectors + [None] + sectors
    tot_a = [sum(t[2][i] for t in trusts) for i in range(5)]
    tot_o = [sum(t[3][i] for t in trusts) for i in range(5)]
    rows = [
        [],
        [None, "Title:", ("Average daily number of available and occupied beds open "
                          "overnight by sector (SYNTHETIC)")],
        [None, "Summary:", ("KH03 is the collection of data to monitor available and "
                            "occupied beds open overnight that are consultant led.")],
        [],
        [None, "Period:", period_text],
        [None, "Source:", "NHS England: SDCS data collection - KH03"],
        [None, "Basis:", "Provider"],
        [None, "Published:", published],
        [None, "Revised:", revised],
        [None, "Status:", "Public"],
        [None, "Contact:", "synthetic@example.org"],
        [],
        [None, "Provider Level Data", None, None, 227],
        [None] * 6 + ["Available"] + [None] * 5 + ["Occupied"] + [None] * 5 + ["% Occupied"],
        header,
        _kh03_block([None, year, month, None, None], "England", tot_a, tot_o),
        [],
    ] + [_kh03_block([None, year, month, t[0], t[1][0]], t[1][1], t[2], t[3])
         for t in trusts]
    dq = [["Data Quality"], ["KH03 synthetic"], []] + dq_rows
    if dq_first:
        _sheet(wb, "Data Quality", dq, first=True)
        _sheet(wb, "NHS Trust by Sector", rows)
    else:
        _sheet(wb, "NHS Trust by Sector", rows, first=True)
    _sheet(wb, "Region by Sector", [[None, "Title:", "synthetic"]])
    _sheet(wb, "Occupied by Specialty", [[None, "Title:", "synthetic"]])
    if not dq_first:
        _sheet(wb, "Data Quality", dq)
    wb.save(path)
    return path


def kh03_published_field() -> Path:
    """2013-14 to 2022-23: blank column A, 'Published:'/'Revised:' rows as text."""
    return _kh03_modern(
        HERE / "kh03_published_field.xlsx", published="18th August 2016",
        revised="23rd November 2017", parent="Region Code", month_col="Period",
        ld="Learning Disabilities", period_text="April to June 2016", year="2016-17",
        month="June", dq_first=False,
        dq_rows=[["This file includes all data submitted up to 1st January 2099."], [],
                 ["Organisations that have revised their data for this quarter:"],
                 ["ZZB", "ANOTHER SYNTHETIC NHS FOUNDATION TRUST"]],
        trusts=[("Y90", ("ZZA", "SYNTHETIC HOSPITALS NHS TRUST"), [410, 370, 0, 40, 0],
                 [350, 330, 0, 20, 0]),
                ("Y91", ("ZZB", "ANOTHER SYNTHETIC NHS FOUNDATION TRUST"),
                 [610, 510, 0, 50, 50], [530, 460, 0, 30, 40])])


def kh03_period_end() -> Path:
    """2023-24 on: Data Quality sheet first, 'Period End', Revised as an Excel date,
    a trust with no G&A beds ('-' occupancy) and an estimate made on a trust's behalf."""
    return _kh03_modern(
        HERE / "kh03_period_end.xlsx", published="23rd May 2024",
        revised=date(2024, 6, 19), parent="Region Code", month_col="Period End",
        ld="Learning Disability", period_text="January to March 2024", year="2023-24",
        month="March", dq_first=True,
        dq_rows=[["This file includes all data submitted up to 1st January 2099."], [],
                 [("The following trusts did not submit and an estimate was submitted on "
                   "their behalf:")],
                 ["ZZC", "SYNTHETIC MENTAL HEALTH NHS TRUST"]],
        trusts=[("Y90", ("ZZA", "SYNTHETIC HOSPITALS NHS TRUST"), [420, 380, 0, 40, 0],
                 [370, 350, 0, 20, 0]),
                ("Y91", ("ZZC", "SYNTHETIC MENTAL HEALTH NHS TRUST"), [90, 0, 0, 0, 90],
                 [80, 0, 0, 0, 80])])


# --------------------------------------------------------------------------------------
# Discharge
# --------------------------------------------------------------------------------------
def _cover(wb: Workbook, start: date, end: date, published, revised) -> None:
    _sheet(wb, "Cover Sheet", [
        [],
        [None, "Title:", "Acute Daily Discharge Situation Report (SYNTHETIC)"],
        [None, "Summary:", "Synthetic fixture."],
        [None, "Period:", start, "-", end],
        [None, "Source:", "NHS England data collection"],
        [None, "Basis:", "Provider"],
        [None, "Published:", published],
        [None, "Revised:", revised],
        [None, "Status:", "Published"],
        [None, "Contact:", "synthetic@example.org"]], first=True)


# Per trust, per day: (nctr, discharged by 17:00, discharged 17:01-23:59); None = '-'.
TRUST_DAYS = {
    ("EAST OF ENGLAND", "ZZA", "SYNTHETIC HOSPITALS NHS TRUST"):
        [(100, 20, 10), (110, 25, 5), (90, 15, 15), (100, 30, 10)],
    ("LONDON", "ZZB", "ANOTHER SYNTHETIC NHS FOUNDATION TRUST"):
        [(50, 10, 0), None, (70, 20, 10), (60, 10, 10)],
}


def discharge_2022_split() -> Path:
    """First 2022 publications: discharges split at 17:00 (four columns per day), helper
    rows of data-item codes and repeated dates, no ICB section, 'All Acute Trusts'."""
    wb = Workbook()
    start = date(2022, 7, 1)
    days = [start + timedelta(days=i) for i in range(4)]
    _cover(wb, start, date(2022, 7, 31), date(2022, 8, 11), "-")
    _sheet(wb, "Contents", [[2.0], [None, "Contents"]])
    _sheet(wb, "Table 1", [[None, "Table 1: Number of organisations that did not submit"]])
    codes, rep, starts, items, names = [None, "Contents", None, None], \
        [None, "Note: Due to disclosure risks, trust-level data exclude hospice discharges.",
         None, None], [None] * 4, [None] * 4, [None] * 4
    for d in days:
        codes += [f"{d:%y%m%d}DIS002_TOTAL", "DIS003_TOTAL", "DIS003b_TOTAL", None]
        rep += [d, d, d, None]
        starts += [d, None, None, None]
        items += ["DIS002_TOTAL", "DIS003_TOTAL", "DIS003b_TOTAL", "DIS004_TOTAL"]
        names += ["Number of patients who no longer meet the criteria to reside",
                  "Number of patients discharged by 17:00",
                  "Number of patients discharged between 17:01 and 23:59",
                  ("Number of patients remaining in hospital who no longer meet the "
                   "criteria to reside")]

    def cells(vals):
        out = []
        for v in vals:
            out += ["-"] * 4 if v is None else [v[0], v[1], v[2], v[0] - v[1] - v[2]]
        return out

    eng = [tuple(sum((t[i] or (0, 0, 0))[k] for t in TRUST_DAYS.values()) for k in range(3))
           for i in range(len(days))]
    rows = [[None, ("Table 2: Number of patients who no longer meet criteria to reside and "
                    "reasons (SYNTHETIC)")], codes, rep, starts, items, names,
            [None, None, None, "ENGLAND (All Acute Trusts)"] + cells(eng),
            [None, None, None, "EAST OF ENGLAND"] + cells(next(iter(TRUST_DAYS.values()))),
            [], [None, "Region", "Org Code", "Org Name"]]
    rows += [[None, *k] + cells(v) for k, v in TRUST_DAYS.items()]
    rows += [[], [None, ("Note: No data is available for 2nd July 2022 due to a data "
                         "collection error (SYNTHETIC).")]]
    _sheet(wb, "Table 2", rows)
    _sheet(wb, "Table 3", [[None, "Table 3: synthetic"]])
    out = HERE / "discharge_2022_split.xlsx"
    wb.save(out)
    return out


def discharge_2024_single() -> Path:
    """Since the 2023 re-issues: three columns per day, an ICB section above the trust
    section, 'Type 1 Trusts', seven tables (from the 2024 specification)."""
    wb = Workbook()
    start = date(2024, 8, 1)
    days = [start + timedelta(days=i) for i in range(4)]
    _cover(wb, start, date(2024, 8, 31), date(2024, 9, 12), None)
    _sheet(wb, "Contents", [[2.0], [None, "Contents"]])
    _sheet(wb, "Table 1", [[None, "Table 1: Number of organisations for whom no data"]])
    dates, names = [None] * 4, [None, "Notes: ", None, None]
    for d in days:
        dates += [d, None, None]
        names += ["Number of patients who no longer meet the criteria to reside",
                  "Number of patients discharged",
                  ("Number of patients remaining in hospital who no longer meet the "
                   "criteria to reside")]

    def cells(vals):
        out = []
        for v in vals:
            out += ["-"] * 3 if v is None else [v[0], v[1] + v[2], v[0] - v[1] - v[2]]
        return out

    eng = [tuple(sum((t[i] or (0, 0, 0))[k] for t in TRUST_DAYS.values()) + (5 if k == 0 else 0)
                 for k in range(3)) for i in range(len(days))]
    rows = [[None, ("Table 2: Number of patients who no longer meet criteria to reside and "
                    "number of patients who were / were not discharged (SYNTHETIC)")],
            [None, "Contents"], dates, names,
            [None, None, None, "ENGLAND (Type 1 Trusts)"] + cells(eng),
            [None, None, None, "EAST OF ENGLAND"] + cells(next(iter(TRUST_DAYS.values()))),
            [], [None, "Region", "ICB Code", "ICB Name"],
            [None, "EAST OF ENGLAND", "ZZ9", "NHS SYNTHETIC INTEGRATED CARE BOARD"]
            + cells(next(iter(TRUST_DAYS.values()))),
            [], [None, "Region", "Org Code", "Org Name"]]
    rows += [[None, *k] + cells(v) for k, v in TRUST_DAYS.items()]
    _sheet(wb, "Table 2", rows)
    for n in range(3, 8):
        _sheet(wb, f"Table {n}", [[None, f"Table {n}: synthetic"]])
    out = HERE / "discharge_2024_single.xlsx"
    wb.save(out)
    return out


if __name__ == "__main__":
    for fn in (kh03_status_line, kh03_published_field, kh03_period_end, discharge_2022_split,
               discharge_2024_single):
        print(fn())
